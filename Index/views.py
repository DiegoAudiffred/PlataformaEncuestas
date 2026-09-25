import email
from email.message import EmailMessage
from email.mime.text import MIMEText
import json
import smtplib

from django.contrib import messages
from django.db.models import Q, Avg, Max, Min, OuterRef, Subquery
from django.db.models.aggregates import Count
from django.db.models.functions import Coalesce
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib.auth.decorators import user_passes_test,login_required
from Index.forms import *
from db.models import * 
from django.core.paginator import Paginator
from django.core.exceptions import PermissionDenied
from django.db import transaction



@login_required(login_url='/login/')    

def index(request):

    
    return render(request, 'index/index.html')


@login_required(login_url='/login/')
def indexEncuestas(request):
    usuario = request.user

    subconsulta_intentos = RespuestaEncuesta.objects.filter(
        encuesta_id=OuterRef('pk'),
        empleado_id=usuario.id
    ).values('encuesta').annotate(
        total=Count('id')
    ).values('total')

    encuestas = Encuesta.objects.filter(
        activa=True,
        dirigido=usuario
    ).annotate(
        intentos_realizados=Coalesce(Subquery(subconsulta_intentos), 0)
    ).order_by('-id')
    creadoes = User.objects.all().order_by('username')

    paginator = Paginator(encuestas, 10)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    context = {
        'encuestas': page_obj,
        'creadoes': creadoes,
    }

    return render(request, 'Index/indexEncuestas.html', context)

@login_required(login_url='/login/')
def encuestasUsuario(request):
    encuestas = Encuesta.objects.filter(creador=request.user).order_by('-id')
    creadoes = User.objects.all().order_by('username')

    paginator = Paginator(encuestas, 10)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    context = {
        'encuestas': page_obj,
        'creadoes': creadoes,
    }

    return render(request, 'Index/misEncuestasUsuario.html', context)

@login_required(login_url='/login/')
def buscarEncuestas(request):
    usuario = request.user
    titulo = request.GET.get('nombre', '').strip()
    estado = request.GET.get('estado', '0').strip()
    autor = request.GET.get('autor', '0').strip()
    page = request.GET.get('page', 1)

    subconsulta_intentos = RespuestaEncuesta.objects.filter(
        encuesta_id=OuterRef('pk'),
        empleado_id=usuario.id
    ).values('encuesta').annotate(
        total=Count('id')
    ).values('total')

    encuestas = Encuesta.objects.filter(
        activa=True,
        dirigido=usuario
    ).annotate(
        intentos_realizados=Coalesce(Subquery(subconsulta_intentos), 0)
    ).order_by('-id')

    if titulo:
        encuestas = encuestas.filter(
            titulo__icontains=titulo
        )

    if estado == '1':
        encuestas = encuestas.filter(
            activa=True
        )
    elif estado == '2':
        encuestas = encuestas.filter(
            activa=False
        )

    if autor and autor != '0':
        encuestas = encuestas.filter(
            creador_id=autor
        )

    paginator = Paginator(encuestas, 10)
    page_obj = paginator.get_page(page)

    contexto = {
        'encuestas': page_obj
    }

    return render(
        request,
        'Index/listaEncuestas.html',
        contexto
    )
@login_required(login_url='/login/')
def buscarEncuestasSelfUser(request):
    titulo = request.GET.get('nombre', '').strip()
    estado = request.GET.get('estado', '0').strip()
    autor = request.GET.get('autor', '0').strip()
    page = request.GET.get('page', 1)

    encuestas = Encuesta.objects.all().order_by('-id')

    if titulo:
        encuestas = encuestas.filter(
            titulo__icontains=titulo
        )

    if estado == '1':
        encuestas = encuestas.filter(
            activa=True
        )
    elif estado == '2':
        encuestas = encuestas.filter(
            activa=False
        )

    if autor and autor != '0':
        encuestas = encuestas.filter(
            creador_id=autor
        )

    paginator = Paginator(encuestas, 10)
    page_obj = paginator.get_page(page)

    contexto = {
        'encuestas': page_obj
    }

    return render(
        request,
        'Index/listaEncuestasSelfUser.html',
        contexto
    )


@login_required(login_url='/login/')
def encuestasContestadas(request):
    """
    Encuestas en las que el usuario actual ya registró al menos una
    respuesta, con el estado de su último intento (completo o no).
    """

    respuestas_qs = (
        RespuestaEncuesta.objects
        .filter(empleado=request.user)
        .select_related('encuesta')
        .prefetch_related('detalles')
        .order_by('encuesta_id', '-intento')
    )

    resumen_por_encuesta = {}

    for r in respuestas_qs:

        if r.encuesta_id not in resumen_por_encuesta:

            total_preguntas = r.encuesta.preguntas.count()
            preguntas_respondidas = (
                r.detalles.values('pregunta_id').distinct().count()
            )

            resumen_por_encuesta[r.encuesta_id] = {
                'encuesta': r.encuesta,
                'intentos_realizados': 0,
                'ultimo_intento': r,
                'completa': (
                    total_preguntas > 0
                    and preguntas_respondidas >= total_preguntas
                ),
            }

        resumen_por_encuesta[r.encuesta_id]['intentos_realizados'] += 1

    items = sorted(
        resumen_por_encuesta.values(),
        key=lambda item: item['encuesta'].id,
        reverse=True
    )

    paginator = Paginator(items, 10)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    return render(
        request,
        'Index/encuestasContestadas.html',
        {'items': page_obj}
    )


@login_required(login_url='/login/')
def verRespuestasEncuesta(request, id):
    """
    Detalle de TODOS los intentos que el usuario actual hizo sobre una
    encuesta puntual: qué contestó en cada pregunta, intento por intento.
    """

    encuesta = get_object_or_404(Encuesta, id=id)

    respuestas = (
        RespuestaEncuesta.objects
        .filter(encuesta=encuesta, empleado=request.user)
        .prefetch_related(
            'detalles__pregunta',
            'detalles__opcion_seleccionada',
            'detalles__usuarios_seleccionados',
        )
        .order_by('-intento')
    )

    if not respuestas.exists():
        raise PermissionDenied('No tienes respuestas registradas en esta encuesta.')

    preguntas = list(encuesta.preguntas.all())
    total_preguntas = len(preguntas)

    intentos = []

    for r in respuestas:

        detalles_por_pregunta = {
            d.pregunta_id: d for d in r.detalles.all()
        }

        preguntas_resumen = []

        for pregunta in preguntas:
            preguntas_resumen.append({
                'pregunta': pregunta,
                'detalle': detalles_por_pregunta.get(pregunta.id),
            })

        intentos.append({
            'respuesta': r,
            'preguntas': preguntas_resumen,
            'completa': (
                total_preguntas > 0
                and len(detalles_por_pregunta) >= total_preguntas
            ),
        })

    return render(
        request,
        'Index/misRespuestas.html',
        {
            'encuesta': encuesta,
            'intentos': intentos,
        }
    )


@login_required(login_url='/login/')
def crearEncuesta(request):
    if request.method == 'POST':
        form = crearEncuestaForm(request.POST, request.FILES)
        if form.is_valid():
            encuesta = form.save(commit=False)
            encuesta.activa = True
            encuesta.creador = request.user
            encuesta.save()
            
            dirigido_seleccionados = form.cleaned_data.get('dirigido')
            if not dirigido_seleccionados:
                usuarios_activos = User.objects.filter(is_active=True)
                encuesta.dirigido.set(usuarios_activos)
            else:
                form.save_m2m()
                
            return redirect('Index:editarEncuesta', id=encuesta.id)
        else:
            print(form.errors)
    else:
        form = crearEncuestaForm()

    return render(request, 'index/crearEncuestas.html', {'form': form})


@login_required(login_url='/login/')
def editarEncuesta(request, id):

    encuesta = get_object_or_404(Encuesta, id=id)

    # --- PERMISOS: solo el creador o staff puede tocar esta encuesta ---
    if not encuesta.puede_editar(request.user):
        raise PermissionDenied('No tienes permiso para editar esta encuesta.')

    if request.method == 'POST':

        action = request.POST.get('action')

        # ================= EDITAR ENCUESTA =================
        if action == 'editar_encuesta':

            form_encuesta = crearEncuestaForm(
                request.POST,
                request.FILES,
                instance=encuesta
            )

            if form_encuesta.is_valid():

                with transaction.atomic():
                    encuesta = form_encuesta.save()

                return JsonResponse({
                    'status': 'success',
                    'encuesta': {
                        'id': encuesta.id,
                        'titulo': encuesta.titulo,
                        'descripcion': encuesta.descripcion,
                        'activa': encuesta.activa,
                        'fechaVencimiento': (
                            encuesta.fechaVencimiento.strftime('%Y-%m-%d')
                            if encuesta.fechaVencimiento
                            else ''
                        ),
                        'imagen_url': (
                            encuesta.imagen.url
                            if encuesta.imagen
                            else None
                        )
                    }
                })

            return JsonResponse({
                'status': 'error',
                'errors': form_encuesta.errors
            }, status=400)

        # ================= ELIMINAR PREGUNTA =================
        if action == 'delete':

            pregunta_id = request.POST.get('pregunta_id')

            pregunta = get_object_or_404(
                Pregunta,
                id=pregunta_id,
                encuesta=encuesta
            )

            tiene_respuestas = DetalleRespuesta.objects.filter(
                pregunta=pregunta
            ).exists()

            if tiene_respuestas and request.POST.get('force') != '1':
                return JsonResponse({
                    'status': 'confirm_required',
                    'message': (
                        'Esta pregunta ya tiene respuestas registradas de '
                        'empleados. Si la eliminas, se perderá ese historial. '
                        '¿Deseas continuar?'
                    )
                }, status=409)

            with transaction.atomic():
                pregunta.delete()

            return JsonResponse({'status': 'success'})

        # ================= CREAR / EDITAR PREGUNTA =================

        pregunta_id = request.POST.get('pregunta_id')

        if pregunta_id:

            pregunta = get_object_or_404(
                Pregunta,
                id=pregunta_id,
                encuesta=encuesta
            )

            form = PreguntaForm(
                request.POST,
                request.FILES,
                instance=pregunta
            )

        else:

            form = PreguntaForm(
                request.POST,
                request.FILES
            )

        if not form.is_valid():
            return JsonResponse({
                'status': 'error',
                'errors': form.errors
            }, status=400)

        tipo = form.cleaned_data['tipo']

        # --- Validación de datos según el tipo ANTES de guardar nada ---
        opciones_list, puntos_list, correctas_list, ids_list = [], [], [], []
        usuarios_ids = []

        if tipo == 'OPCION_MULTIPLE':

            opciones_list = [
                t.strip() for t in request.POST.getlist('opciones[]')
            ]
            puntos_list = request.POST.getlist('puntos[]')
            correctas_list = request.POST.getlist('es_correcta[]')
            ids_list = request.POST.getlist('opcion_id[]')

            if not any(opciones_list):
                return JsonResponse({
                    'status': 'error',
                    'errors': {'opciones': ['Agrega al menos una opción.']}
                }, status=400)

        elif tipo in Pregunta.TIPOS_USUARIO:

            usuarios_ids = request.POST.getlist('usuarios_opciones[]')

            if not usuarios_ids:
                return JsonResponse({
                    'status': 'error',
                    'errors': {
                        'usuarios_opciones': [
                            'Selecciona al menos un usuario disponible '
                            'para esta pregunta.'
                        ]
                    }
                }, status=400)

        # --- Guardado atómico ---
        with transaction.atomic():

            pregunta = form.save(commit=False)
            pregunta.encuesta = encuesta
            pregunta.save()

            # Limpieza de relaciones que no aplican al tipo actual
            if tipo != 'OPCION_MULTIPLE':
                pregunta.opciones.all().delete()
            if tipo not in Pregunta.TIPOS_USUARIO:
                pregunta.usuarios_opciones.clear()

            opciones_creadas = []

            if tipo == 'OPCION_MULTIPLE':

                ids_vistos = set()

                for index, texto in enumerate(opciones_list):

                    texto_clean = texto.strip()
                    if not texto_clean:
                        continue

                    pts = 0
                    if (
                        index < len(puntos_list)
                        and puntos_list[index].strip().lstrip('-').isdigit()
                    ):
                        pts = int(puntos_list[index])

                    es_cor = str(index) in correctas_list

                    opcion_id = (
                        ids_list[index]
                        if index < len(ids_list) and ids_list[index]
                        else None
                    )

                    if opcion_id:
                        # Actualiza la opción existente en vez de recrearla,
                        # así no se rompen respuestas ya guardadas que
                        # apuntan a este OpcionRespuesta.id
                        opcion = get_object_or_404(
                            OpcionRespuesta, id=opcion_id, pregunta=pregunta
                        )
                        opcion.texto_opcion = texto_clean
                        opcion.puntos = pts
                        opcion.es_correcta = es_cor
                        opcion.orden = index
                        opcion.save()
                        ids_vistos.add(opcion.id)
                    else:
                        opcion = OpcionRespuesta.objects.create(
                            pregunta=pregunta,
                            texto_opcion=texto_clean,
                            puntos=pts,
                            es_correcta=es_cor,
                            orden=index,
                        )
                        ids_vistos.add(opcion.id)

                    opciones_creadas.append({
                        'id': opcion.id,
                        'texto': opcion.texto_opcion,
                        'puntos': opcion.puntos,
                        'es_correcta': opcion.es_correcta,
                    })

                # Borra en BD las opciones que el usuario quitó del form
                pregunta.opciones.exclude(id__in=ids_vistos).delete()

            elif tipo in Pregunta.TIPOS_USUARIO:

                usuarios_validos = User.objects.filter(id__in=usuarios_ids)
                pregunta.usuarios_opciones.set(usuarios_validos)

                # 'opciones_creadas' se deja vacío a propósito: la lista de
                # usuarios va SOLO en 'usuarios_opciones_payload' más abajo.
                # Antes se llenaban ambas listas con los mismos usuarios y
                # el frontend los pintaba dos veces en la tarjeta.
                opciones_creadas = []

            image_url = pregunta.imagen.url if pregunta.imagen else None

            usuarios_opciones_payload = (
                [
                    {'id': u.id, 'nombre': u.nombre or u.username}
                    for u in pregunta.usuarios_opciones.all()
                ]
                if pregunta.es_tipo_usuario()
                else []
            )

        return JsonResponse({
            'status': 'success',
            'is_edit': bool(pregunta_id),
            'pregunta': {
                'id': pregunta.id,
                'texto': pregunta.texto_pregunta,
                'tipo_val': pregunta.tipo,
                'tipo': pregunta.get_tipo_display(),
                'imagen_url': image_url,
                'opciones': opciones_creadas,
                'usuarios_opciones': usuarios_opciones_payload,
            }
        })

    # ================= GET =================
    form = PreguntaForm()
    form_encuesta = crearEncuestaForm(instance=encuesta)

    preguntas = (
        encuesta.preguntas
        .prefetch_related('opciones', 'usuarios_opciones')
        .all()
    )

    # Universo de usuarios que se puede ofrecer como opción de respuesta al
    # crear una pregunta tipo "Selección de Usuario": TODOS los usuarios
    # activos del sistema, no solo los dirigidos a esta encuesta (quien
    # responde puede tener que elegir a cualquier compañero, no solo a
    # quienes también fueron invitados a contestar).
    usuarios_disponibles = User.objects.filter(is_active=True).order_by('username')

    return render(
        request,
        'index/editarEncuesta.html',
        {
            'encuesta': encuesta,
            'preguntas': preguntas,
            'form': form,
            'form_encuesta': form_encuesta,
            'usuarios_disponibles': usuarios_disponibles,
        }
    )

@login_required(login_url='/login/')    
def encuestasResultados(request):
    return render(request, 'index/encuestasResultados.html')

@login_required(login_url='/login/')
def responderEncuesta(request, id):

    encuesta = get_object_or_404(
        Encuesta,
        id=id
    )

    if not encuesta.activa:
        return render(
            request,
            'index/responderEncuesta.html',
            {
                'encuesta': encuesta,
                'error': 'Esta encuesta ya no está activa.'
            }
        )

    if encuesta.fechaVencimiento:
        if timezone.localdate() > encuesta.fechaVencimiento:
            return render(
                request,
                'index/responderEncuesta.html',
                {
                    'encuesta': encuesta,
                    'error': 'Esta encuesta ya venció.'
                }
            )

    if not encuesta.dirigido.filter(
        id=request.user.id
    ).exists():
        return render(
            request,
            'index/responderEncuesta.html',
            {
                'encuesta': encuesta,
                'error': 'No tienes autorización para responder esta encuesta.'
            }
        )

    intentos_realizados = RespuestaEncuesta.objects.filter(
        encuesta=encuesta,
        empleado=request.user
    ).count()

    intentos_restantes = (
        encuesta.maxIntentos - intentos_realizados
    )

    if intentos_restantes <= 0:

        ultima_respuesta = (
            RespuestaEncuesta.objects
            .filter(
                encuesta=encuesta,
                empleado=request.user
            )
            .order_by('-intento')
            .first()
        )

        return render(
            request,
            'index/responderEncuesta.html',
            {
                'encuesta': encuesta,
                'bloqueada': True,
                'intentos_realizados': intentos_realizados,
                'intentos_restantes': 0,
                'ultima_respuesta': ultima_respuesta
            }
        )

    preguntas = (
        encuesta.preguntas
        .prefetch_related('opciones', 'usuarios_opciones')
        .all()
    )

    if request.method == 'POST':

        intentos_realizados = RespuestaEncuesta.objects.filter(
            encuesta=encuesta,
            empleado=request.user
        ).count()

        if intentos_realizados >= encuesta.maxIntentos:

            return JsonResponse({
                'status': 'error',
                'mensaje': 'Ya alcanzaste el número máximo de intentos.'
            }, status=400)

        intento = intentos_realizados + 1

        respuesta_encuesta = RespuestaEncuesta.objects.create(
            encuesta=encuesta,
            empleado=request.user,
            intento=intento,
            puntuacion=0
        )

        puntuacion_total = 0

        for pregunta in preguntas:

            if pregunta.tipo == 'USUARIO_MULTIPLE':

                usuario_ids = request.POST.getlist(
                    f'pregunta_{pregunta.id}[]'
                )

                if not usuario_ids:
                    continue

                detalle = DetalleRespuesta.objects.create(
                    respuesta_encuesta=respuesta_encuesta,
                    pregunta=pregunta
                )

                usuarios_validos = pregunta.usuarios_opciones.filter(
                    id__in=usuario_ids
                )

                detalle.usuarios_seleccionados.set(usuarios_validos)

                continue

            respuesta = request.POST.get(
                f'pregunta_{pregunta.id}'
            )

            if not respuesta:
                continue

            if pregunta.tipo == 'OPCION_MULTIPLE':

                opcion = get_object_or_404(
                    OpcionRespuesta,
                    id=respuesta,
                    pregunta=pregunta
                )

                DetalleRespuesta.objects.create(
                    respuesta_encuesta=respuesta_encuesta,
                    pregunta=pregunta,
                    opcion_seleccionada=opcion
                )

                if opcion.es_correcta:
                    puntuacion_total += opcion.puntos

            elif pregunta.tipo == 'TEXTO':

                DetalleRespuesta.objects.create(
                    respuesta_encuesta=respuesta_encuesta,
                    pregunta=pregunta,
                    texto_respuesta=respuesta
                )

            elif pregunta.tipo == 'USUARIO_UNICO':

                usuario = pregunta.usuarios_opciones.filter(
                    id=respuesta
                ).first()

                if not usuario:
                    continue

                detalle = DetalleRespuesta.objects.create(
                    respuesta_encuesta=respuesta_encuesta,
                    pregunta=pregunta
                )

                detalle.usuarios_seleccionados.add(usuario)

        respuesta_encuesta.puntuacion = puntuacion_total
        respuesta_encuesta.save(update_fields=['puntuacion'])

        return JsonResponse({
            'status': 'success',
            'mensaje': 'Encuesta respondida correctamente.',
            'intento': intento,
            'puntuacion': puntuacion_total,
            'intentos_restantes': encuesta.maxIntentos - intento
        })

    return render(
        request,
        'index/responderEncuesta.html',
        {
            'encuesta': encuesta,
            'preguntas': preguntas,
            'intentos_realizados': intentos_realizados,
            'intentos_restantes': intentos_restantes,
            'bloqueada': False
        }
    )

@login_required(login_url='/login/')
def enviarCorreoEncuesta(request, encuesta_id):
    if request.method == 'POST':
        encuesta = get_object_or_404(Encuesta, id=encuesta_id)
        
        destinatarios = list(encuesta.dirigido.filter(is_active=True).exclude(email='').values_list('email', flat=True))
        dominio = "http://192.168.0.29:8001/responderEncuesta/"
        url_final = f"{dominio}{encuesta.id}"

        if destinatarios:
            host = "ucg.com.mx"
            port = 587
            usuario_smtp = "informacion@ucg.com.mx"
            password_smtp = "UcG911_@!#"

            asunto = f"Nueva encuesta disponible: {encuesta.titulo}"

            cuerpo_html = f"""
            <html>
                <body style="font-family: Arial, sans-serif; color: #333; line-height: 1.6;">
                    <div style="max-width: 600px; margin: 0 auto; border: 1px solid #ddd; padding: 20px; border-radius: 10px;">
                        <h2 style="color: #2c3e50;">Nueva Encuesta Asignada</h2>
                        <p>Estimado usuario,</p>
                        <p>Se le ha asignado la encuesta <strong>{encuesta.titulo}</strong> creada por {request.user.get_full_name() or request.user.username}.</p>
                        <p>Le invitamos a responderla ingresando al siguiente enlace:</p>
                        <div style="margin: 30px 0; text-align: center;">
                            <a href="{url_final}" 
                               style="background-color: #007bff; color: white; padding: 12px 25px; text-decoration: none; border-radius: 5px; font-weight: bold; display: inline-block;">
                                Responder Encuesta
                            </a>
                        </div>
                        <p style="font-size: 12px; color: #777;">Si el botón no funciona, copie y pegue el siguiente enlace en su navegador:<br>{url_final}</p>
                    </div>
                </body>
            </html>
            """

            mensaje = EmailMessage()
            mensaje['Subject'] = asunto
            mensaje['From'] = usuario_smtp
            mensaje['To'] = ", ".join(destinatarios)
            mensaje.add_alternative(cuerpo_html, subtype="html")

            try:
                server = smtplib.SMTP(host, port)
                server.starttls()
                server.login(usuario_smtp, password_smtp)
                server.sendmail(usuario_smtp, destinatarios, mensaje.as_string())
                server.quit()
                messages.success(request, "Correos enviados exitosamente.")
            except Exception as e:
                messages.error(request, f"Error al enviar los correos: {e}")
        else:
            messages.warning(request, "No hay usuarios con correo válido para notificar.")

        return redirect('Index:editarEncuesta', id=encuesta.id)

    return redirect('Index:indexEncuestas')


"""
Vistas para la sección "Resultados de mis encuestas".

Requiere que ya tengas importado en tu views.py:
    from django.db.models import Avg, Max, Min, Q
    from django.db.models.aggregates import Count
    from django.http import HttpResponse
    from django.core.exceptions import PermissionDenied
    from django.shortcuts import get_object_or_404, render

Agrega también, si no los tienes ya:
    import json
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

openpyxl no viene en Django por defecto: si no lo tienes instalado,
`pip install openpyxl`.

Usa los modelos Encuesta, Pregunta, OpcionRespuesta, RespuestaEncuesta,
DetalleRespuesta y User — ya deberían estar disponibles en tu views.py via
`from db.models import *`.

Rutas a agregar en urls.py:
    path("encuestasResultados/", views.encuestasResultados, name='encuestasResultados'),
    path("encuestasResultados/<int:id>/", views.encuestaResultadoDetalle, name='encuestaResultadoDetalle'),
    path("encuestasResultados/<int:id>/excel/", views.exportarResultadosExcel, name='exportarResultadosExcel'),
    path("respuestaParticipante/<int:respuesta_id>/", views.verRespuestaParticipante, name='verRespuestaParticipante'),
"""


@login_required(login_url='/login/')
def encuestasResultados(request):
    """
    Encuestas creadas por el usuario actual (o todas, si es staff/superuser),
    con un resumen de participación para decidir a cuál entrar a analizar.
    """

    if request.user.is_staff or request.user.is_superuser:
        encuestas = Encuesta.objects.all()
    else:
        encuestas = Encuesta.objects.filter(creador=request.user)

    encuestas = encuestas.order_by('-id')

    items = []

    for encuesta in encuestas:

        respuestas = RespuestaEncuesta.objects.filter(encuesta=encuesta)

        agg = respuestas.aggregate(promedio=Avg('puntuacion'))

        items.append({
            'encuesta': encuesta,
            'participantes': respuestas.values('empleado').distinct().count(),
            'intentos_totales': respuestas.count(),
            'total_preguntas': encuesta.preguntas.count(),
            'promedio': (
                round(agg['promedio'], 1)
                if agg['promedio'] is not None
                else None
            ),
            'tiene_puntaje': OpcionRespuesta.objects.filter(
                pregunta__encuesta=encuesta, puntos__gt=0
            ).exists(),
        })

    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(
        request,
        'Index/encuestasResultados.html',
        {'items': page_obj}
    )


@login_required(login_url='/login/')
def encuestaResultadoDetalle(request, id):
    """
    Panel de análisis de una encuesta: participación, distribución de
    respuestas por pregunta (incluye ranking de "más elegido" para las
    preguntas de usuario) y puntuaciones, si la encuesta usa puntos.
    """

    encuesta = get_object_or_404(Encuesta, id=id)

    if not encuesta.puede_editar(request.user):
        raise PermissionDenied('No tienes permiso para ver los resultados de esta encuesta.')

    respuestas = (
        RespuestaEncuesta.objects
        .filter(encuesta=encuesta)
        .select_related('empleado')
        .order_by('-fecha_respuesta')
    )

    total_participantes = respuestas.values('empleado').distinct().count()
    agg_puntaje = respuestas.aggregate(
        promedio=Avg('puntuacion'), maximo=Max('puntuacion'), minimo=Min('puntuacion')
    )

    tiene_puntaje = OpcionRespuesta.objects.filter(
        pregunta__encuesta=encuesta, puntos__gt=0
    ).exists()

    preguntas = encuesta.preguntas.prefetch_related('opciones', 'usuarios_opciones').all()

    analitica_preguntas = []

    for pregunta in preguntas:

        total_resp_pregunta = DetalleRespuesta.objects.filter(pregunta=pregunta).count()

        if pregunta.tipo == 'OPCION_MULTIPLE':

            data = []
            for opcion in pregunta.opciones.all():
                count = DetalleRespuesta.objects.filter(opcion_seleccionada=opcion).count()
                pct = round(count / total_resp_pregunta * 100, 1) if total_resp_pregunta else 0
                data.append({
                    'etiqueta': opcion.texto_opcion,
                    'puntos': opcion.puntos,
                    'es_correcta': opcion.es_correcta,
                    'count': count,
                    'pct': pct,
                })

            analitica_preguntas.append({
                'pregunta': pregunta, 'tipo': 'opciones',
                'data': data, 'total': total_resp_pregunta,
            })

        elif pregunta.tipo in ('USUARIO_UNICO', 'USUARIO_MULTIPLE'):

            ranking = (
                User.objects
                .filter(seleccionado_en_respuestas__pregunta=pregunta)
                .annotate(votos=Count(
                    'seleccionado_en_respuestas',
                    filter=Q(seleccionado_en_respuestas__pregunta=pregunta)
                ))
                .order_by('-votos')
            )

            data = [
                {
                    'etiqueta': u.nombre or u.username,
                    'count': u.votos,
                    'pct': round(u.votos / total_resp_pregunta * 100, 1) if total_resp_pregunta else 0,
                }
                for u in ranking
            ]

            analitica_preguntas.append({
                'pregunta': pregunta, 'tipo': 'usuarios',
                'data': data, 'total': total_resp_pregunta,
            })

        else:  # TEXTO

            respuestas_texto = (
                DetalleRespuesta.objects
                .filter(pregunta=pregunta)
                .exclude(texto_respuesta__isnull=True)
                .exclude(texto_respuesta__exact='')
                .select_related('respuesta_encuesta__empleado')
                .order_by('-respuesta_encuesta__fecha_respuesta')
            )

            analitica_preguntas.append({
                'pregunta': pregunta, 'tipo': 'texto',
                'data': respuestas_texto, 'total': total_resp_pregunta,
            })

    return render(
        request,
        'Index/encuestaResultadoDetalle.html',
        {
            'encuesta': encuesta,
            'respuestas': respuestas,
            'total_participantes': total_participantes,
            'agg_puntaje': agg_puntaje,
            'tiene_puntaje': tiene_puntaje,
            'analitica_preguntas': analitica_preguntas,
            'analitica_preguntas_json': json.dumps([
                {
                    'id': item['pregunta'].id,
                    'labels': [fila['etiqueta'] for fila in item['data']],
                    'valores': [fila['count'] for fila in item['data']],
                }
                for item in analitica_preguntas
                if item['tipo'] in ('opciones', 'usuarios')
            ]),
        }
    )


@login_required(login_url='/login/')
def verRespuestaParticipante(request, respuesta_id):
    """
    El detalle de UN intento específico de UN participante, para que el
    creador de la encuesta (o staff) vea exactamente qué contestó.
    """

    respuesta = get_object_or_404(
        RespuestaEncuesta.objects.select_related('encuesta', 'empleado'),
        id=respuesta_id
    )

    encuesta = respuesta.encuesta

    if not encuesta.puede_editar(request.user):
        raise PermissionDenied('No tienes permiso para ver esta respuesta.')

    detalles_por_pregunta = {
        d.pregunta_id: d
        for d in respuesta.detalles
            .select_related('pregunta', 'opcion_seleccionada')
            .prefetch_related('usuarios_seleccionados')
    }

    preguntas_resumen = []

    for pregunta in encuesta.preguntas.all():
        preguntas_resumen.append({
            'pregunta': pregunta,
            'detalle': detalles_por_pregunta.get(pregunta.id),
        })

    return render(
        request,
        'Index/verRespuestaParticipante.html',
        {
            'encuesta': encuesta,
            'respuesta': respuesta,
            'preguntas': preguntas_resumen,
        }
    )


@login_required(login_url='/login/')
def exportarResultadosExcel(request, id):
    """
    Descarga en .xlsx los resultados de una encuesta: resumen, distribución
    por pregunta (incluye ranking de usuarios más elegidos), respuestas de
    texto libre, y un detalle "crudo" (una fila por pregunta x intento) para
    que se pueda hacer una tabla dinámica en Excel si se quiere ir más a
    fondo.
    """

    encuesta = get_object_or_404(Encuesta, id=id)

    if not encuesta.puede_editar(request.user):
        raise PermissionDenied('No tienes permiso para exportar esta encuesta.')

    respuestas = (
        RespuestaEncuesta.objects
        .filter(encuesta=encuesta)
        .select_related('empleado')
        .order_by('empleado__username', 'intento')
    )

    preguntas = list(encuesta.preguntas.all())

    wb = openpyxl.Workbook()

    font_normal = Font(name='Arial', size=10)
    font_header = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    fill_header = PatternFill('solid', fgColor='2F5597')

    def estilizar_encabezados(ws, fila=1):
        for cell in ws[fila]:
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def aplicar_fuente(ws, min_row=2, wrap=False):
        for row in ws.iter_rows(min_row=min_row):
            for cell in row:
                cell.font = font_normal
                if wrap:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

    def anchos(ws, letras_anchos):
        for col, width in letras_anchos:
            ws.column_dimensions[col].width = width

    # ---------- Hoja: Resumen ----------
    ws = wb.active
    ws.title = 'Resumen'

    agg = respuestas.aggregate(
        promedio=Avg('puntuacion'), maximo=Max('puntuacion'), minimo=Min('puntuacion')
    )
    total_participantes = respuestas.values('empleado').distinct().count()

    ws.append(['Campo', 'Valor'])
    estilizar_encabezados(ws)

    filas_resumen = [
        ('Encuesta', encuesta.titulo),
        ('Descripción', encuesta.descripcion),
        ('Creador', encuesta.creador.username if encuesta.creador else ''),
        ('Activa', 'Sí' if encuesta.activa else 'No'),
        ('Fecha de vencimiento', encuesta.fechaVencimiento.strftime('%d/%m/%Y') if encuesta.fechaVencimiento else 'Sin fecha'),
        ('Máx. intentos permitidos', encuesta.maxIntentos),
        ('Total de preguntas', len(preguntas)),
        ('Participantes que respondieron', total_participantes),
        ('Intentos totales registrados', respuestas.count()),
        ('Puntuación promedio', round(agg['promedio'], 2) if agg['promedio'] is not None else 'N/A'),
        ('Puntuación máxima', agg['maximo'] if agg['maximo'] is not None else 'N/A'),
        ('Puntuación mínima', agg['minimo'] if agg['minimo'] is not None else 'N/A'),
    ]
    for fila in filas_resumen:
        ws.append(list(fila))

    aplicar_fuente(ws)
    anchos(ws, [('A', 28), ('B', 55)])

    # ---------- Hoja: Por Pregunta ----------
    ws2 = wb.create_sheet('Por Pregunta')
    ws2.append(['Pregunta', 'Tipo', 'Opción / Usuario', 'Puntos', 'Veces elegida', '% del total'])
    estilizar_encabezados(ws2)

    for pregunta in preguntas:

        total_resp_pregunta = DetalleRespuesta.objects.filter(pregunta=pregunta).count()

        if pregunta.tipo == 'OPCION_MULTIPLE':

            for opcion in pregunta.opciones.all():
                count = DetalleRespuesta.objects.filter(opcion_seleccionada=opcion).count()
                pct = round(count / total_resp_pregunta * 100, 1) if total_resp_pregunta else 0
                ws2.append([
                    pregunta.texto_pregunta, pregunta.get_tipo_display(),
                    opcion.texto_opcion, opcion.puntos, count, pct
                ])

        elif pregunta.tipo in ('USUARIO_UNICO', 'USUARIO_MULTIPLE'):

            ranking = (
                User.objects
                .filter(seleccionado_en_respuestas__pregunta=pregunta)
                .annotate(votos=Count(
                    'seleccionado_en_respuestas',
                    filter=Q(seleccionado_en_respuestas__pregunta=pregunta)
                ))
                .order_by('-votos')
            )

            for u in ranking:
                pct = round(u.votos / total_resp_pregunta * 100, 1) if total_resp_pregunta else 0
                ws2.append([
                    pregunta.texto_pregunta, pregunta.get_tipo_display(),
                    u.nombre or u.username, '', u.votos, pct
                ])

        else:  # TEXTO
            ws2.append([
                pregunta.texto_pregunta, pregunta.get_tipo_display(),
                f'{total_resp_pregunta} respuesta(s) de texto libre — ver hoja "Texto libre"',
                '', total_resp_pregunta, ''
            ])

    aplicar_fuente(ws2)
    anchos(ws2, [('A', 35), ('B', 22), ('C', 32), ('D', 10), ('E', 14), ('F', 12)])

    # ---------- Hoja: Texto libre ----------
    ws3 = wb.create_sheet('Texto libre')
    ws3.append(['Pregunta', 'Empleado', 'Intento', 'Fecha', 'Respuesta'])
    estilizar_encabezados(ws3)

    for pregunta in preguntas:

        if pregunta.tipo != 'TEXTO':
            continue

        detalles = (
            DetalleRespuesta.objects
            .filter(pregunta=pregunta)
            .exclude(texto_respuesta__isnull=True)
            .exclude(texto_respuesta__exact='')
            .select_related('respuesta_encuesta__empleado')
            .order_by('respuesta_encuesta__empleado__username')
        )

        for d in detalles:
            ws3.append([
                pregunta.texto_pregunta,
                d.respuesta_encuesta.empleado.nombre or d.respuesta_encuesta.empleado.username,
                d.respuesta_encuesta.intento,
                d.respuesta_encuesta.fecha_respuesta.strftime('%d/%m/%Y %H:%M'),
                d.texto_respuesta or '',
            ])

    aplicar_fuente(ws3, wrap=True)
    anchos(ws3, [('A', 30), ('B', 22), ('C', 10), ('D', 18), ('E', 60)])

    # ---------- Hoja: Detalle (una fila por pregunta x intento) ----------
    ws4 = wb.create_sheet('Detalle')
    ws4.append([
        'Empleado', 'Usuario', 'Intento', 'Fecha', 'Puntuación total del intento',
        'Pregunta', 'Tipo', 'Respuesta', 'Puntos de la opción'
    ])
    estilizar_encabezados(ws4)

    for r in respuestas.prefetch_related(
        'detalles__pregunta', 'detalles__opcion_seleccionada', 'detalles__usuarios_seleccionados'
    ):

        detalles_por_pregunta = {d.pregunta_id: d for d in r.detalles.all()}

        for pregunta in preguntas:

            detalle = detalles_por_pregunta.get(pregunta.id)

            if not detalle:
                respuesta_txt = '(sin responder)'
                puntos = ''
            elif pregunta.tipo == 'OPCION_MULTIPLE':
                respuesta_txt = (
                    detalle.opcion_seleccionada.texto_opcion
                    if detalle.opcion_seleccionada
                    else '(opción eliminada)'
                )
                puntos = detalle.opcion_seleccionada.puntos if detalle.opcion_seleccionada else ''
            elif pregunta.tipo == 'TEXTO':
                respuesta_txt = detalle.texto_respuesta or ''
                puntos = ''
            else:
                nombres = [u.nombre or u.username for u in detalle.usuarios_seleccionados.all()]
                respuesta_txt = ', '.join(nombres) if nombres else '(sin responder)'
                puntos = ''

            ws4.append([
                r.empleado.nombre or r.empleado.username,
                r.empleado.username,
                r.intento,
                r.fecha_respuesta.strftime('%d/%m/%Y %H:%M'),
                r.puntuacion,
                pregunta.texto_pregunta,
                pregunta.get_tipo_display(),
                respuesta_txt,
                puntos,
            ])

    aplicar_fuente(ws4)
    anchos(ws4, [
        ('A', 22), ('B', 16), ('C', 8), ('D', 18), ('E', 14),
        ('F', 35), ('G', 20), ('H', 40), ('I', 12),
    ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    nombre_archivo = 'resultados_' + ''.join(
        c if c.isalnum() else '_' for c in encuesta.titulo
    ).strip('_')

    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}.xlsx"'
    wb.save(response)

    return response



def eliminarEncuesta(request,id):
    encuesta = Encuesta.objects.get(id=id)
    encuesta.delete()
    return redirect('Index:encuestasUsuario')